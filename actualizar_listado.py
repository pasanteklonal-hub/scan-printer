import glob
import json
import os
import subprocess
import shutil
from pathlib import Path

import pandas as pd

from openpyxl import load_workbook
from openpyxl.styles import (
    PatternFill,
    Font,
    Border,
    Side,
    Alignment,
)
from openpyxl.utils import get_column_letter

from app_paths import APP_DIR, resource_path


BASE_DIR = APP_DIR
ARCHIVO_BASE = resource_path("impresoras.xlsx")
ARCHIVO_SALIDA = BASE_DIR / "impresoras_al_dia.xlsx"
CARPETA_JSON = BASE_DIR / "estado-impresoras"


def normalizar_mac(mac):

    if pd.isna(mac):
        return None

    return (
        str(mac)
        .strip()
        .upper()
    )


def normalizar_ip(ip):

    if pd.isna(ip):
        return None

    return str(ip).strip()


def obtener_ultimo_json():

    archivos = glob.glob(str(CARPETA_JSON / "impresoras_*.json"))

    if not archivos:
        raise FileNotFoundError(
            "No se encontró ningún archivo impresoras_*.json"
        )

    return max(archivos, key=os.path.getmtime)


def leer_json(ruta_json):

    with open(ruta_json, "r", encoding="utf-8") as f:
        datos = json.load(f)

    indice = {}

    for imp in datos:

        mac = normalizar_mac(imp.get("mac"))
        ip = normalizar_ip(imp.get("ip"))

        if not mac and not ip:
            continue

        toner = ""

        if isinstance(imp.get("toner"), dict):
            if imp["toner"]:
                toner = next(iter(imp["toner"].values()))

        paginas = ""

        if isinstance(imp.get("paginas"), dict):
            paginas = imp["paginas"].get(
                "impresiones_totales",
                "",
            )

        registro = {
            "IP": imp.get("ip", ""),
            "MAC": imp.get("mac") or "",
            "MARCA": imp.get("marca") or "",
            "MODELO": imp.get("modelo") or "",
            "NIVEL_TONER": toner,
            "IMPRESIONES_TOTALES": paginas,
            "NUMERO_SERIE": imp.get("numero_serie", ""),
        }
        if mac:
            indice[mac] = registro
        if ip:
            indice[f"IP:{ip}"] = registro

    return indice


def buscar_impresora(datos, mac, ip):
    ip = normalizar_ip(ip)
    if ip:
        registro = datos.get(f"IP:{ip}")
        if registro:
            return registro

    mac = normalizar_mac(mac)
    return datos.get(mac, {})


def obtener_area_impresora(ip):
    ip = normalizar_ip(ip)
    for ruta in (ARCHIVO_BASE, ARCHIVO_SALIDA):
        if not ruta.exists():
            continue
        wb = load_workbook(ruta, read_only=True, data_only=True)
        ws = wb.active
        columnas = {
            str(c.value).strip().upper(): c.column
            for c in ws[1]
            if c.value is not None
        }
        if "IP" in columnas and "SECTOR" in columnas:
            for fila in ws.iter_rows(min_row=2, values_only=True):
                if normalizar_ip(fila[columnas["IP"] - 1]) == ip:
                    area = fila[columnas["SECTOR"] - 1]
                    wb.close()
                    return str(area).strip() if area else ""
        wb.close()
    return ""


def areas_en_uso():
    areas = set()
    for ruta in (ARCHIVO_BASE, ARCHIVO_SALIDA):
        if not ruta.exists():
            continue
        wb = load_workbook(ruta, read_only=True, data_only=True)
        ws = wb.active
        columnas = {
            str(c.value).strip().upper(): c.column
            for c in ws[1]
            if c.value is not None
        }
        columna_area = columnas.get("SECTOR") or columnas.get("AREA") or columnas.get("ÁREA")
        if columna_area:
            areas.update(
                str(fila[columna_area - 1]).strip()
                for fila in ws.iter_rows(min_row=2, values_only=True)
                if fila[columna_area - 1] is not None
                and str(fila[columna_area - 1]).strip()
            )
        wb.close()
    return areas


def establecer_area_impresora(ip, marca, area):
    ip = normalizar_ip(ip)
    wb = load_workbook(ARCHIVO_BASE)
    ws = wb.active
    columnas = {
        str(c.value).strip().upper(): c.column
        for c in ws[1]
        if c.value is not None
    }
    for campo in ("IP", "MARCA", "SECTOR"):
        if campo not in columnas:
            columnas[campo] = ws.max_column + 1
            ws.cell(row=1, column=columnas[campo], value=campo)

    fila_objetivo = None
    for fila in range(2, ws.max_row + 1):
        if normalizar_ip(ws.cell(fila, columnas["IP"]).value) == ip:
            fila_objetivo = fila
            break

    if fila_objetivo is None:
        mac = None
        if ARCHIVO_SALIDA.exists():
            salida = load_workbook(ARCHIVO_SALIDA, read_only=True, data_only=True)
            hoja_salida = salida.active
            columnas_salida = {
                str(c.value).strip().upper(): c.column
                for c in hoja_salida[1]
                if c.value is not None
            }
            if "IP" in columnas_salida and "MAC" in columnas_salida:
                for fila in hoja_salida.iter_rows(min_row=2, values_only=True):
                    if normalizar_ip(fila[columnas_salida["IP"] - 1]) == ip:
                        mac = normalizar_mac(fila[columnas_salida["MAC"] - 1])
                        break
            salida.close()
        if not mac:
            try:
                mac = leer_json(obtener_ultimo_json()).get(f"IP:{ip}", {}).get("MAC")
            except FileNotFoundError:
                mac = None
        if mac and "MAC" in columnas:
            for fila in range(2, ws.max_row + 1):
                if normalizar_mac(ws.cell(fila, columnas["MAC"]).value) == normalizar_mac(mac):
                    fila_objetivo = fila
                    break

    if fila_objetivo is None:
        fila_objetivo = ws.max_row + 1
        ws.cell(row=fila_objetivo, column=columnas["IP"], value=ip)
        ws.cell(row=fila_objetivo, column=columnas["MARCA"], value=marca)
    else:
        ws.cell(row=fila_objetivo, column=columnas["IP"], value=ip)
        if not ws.cell(fila_objetivo, column=columnas["MARCA"]).value:
            ws.cell(row=fila_objetivo, column=columnas["MARCA"], value=marca)
    ws.cell(row=fila_objetivo, column=columnas["SECTOR"], value=area)
    wb.save(ARCHIVO_BASE)


def aplicar_formato():

    wb = load_workbook(ARCHIVO_SALIDA)
    ws = wb.active

    # -----------------------------
    # Estilos
    # -----------------------------

    azul = PatternFill(
        "solid",
        fgColor="1F4E78"
    )

    verde = PatternFill(
        "solid",
        fgColor="92D050"
    )

    naranja = PatternFill(
        "solid",
        fgColor="FFC000"
    )

    rojo = PatternFill(
        "solid",
        fgColor="FF6666"
    )

    gris = PatternFill(
        "solid",
        fgColor="E7E6E6"
    )

    borde = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    fuente = Font(
        name="Calibri",
        size=10
    )

    fuente_titulo = Font(
        name="Calibri",
        size=11,
        bold=True,
        color="FFFFFF",
    )

    # -----------------------------
    # Encabezados
    # -----------------------------

    for c in ws[1]:
        c.fill = azul
        c.font = fuente_titulo
        c.border = borde
        c.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    # -----------------------------
    # Buscar columnas
    # -----------------------------

    columnas = {}

    for c in ws[1]:
        columnas[c.value] = c.column_letter

    col_ip = columnas["IP"]
    col_toner = columnas["NIVEL_TONER"]
    col_pag = columnas["IMPRESIONES_TOTALES"]
    col_serie = columnas["NUMERO_SERIE"]

    # -----------------------------
    # Filas
    # -----------------------------

    for fila in range(2, ws.max_row + 1):

        ip = ws[f"{col_ip}{fila}"]

        toner = ws[f"{col_toner}{fila}"]

        paginas = ws[f"{col_pag}{fila}"]
        serie = ws[f"{col_serie}{fila}"]

        # Fuente y bordes

        for celda in ws[fila]:
            celda.font = fuente
            celda.border = borde

        ip.alignment = Alignment(horizontal="center")
        toner.alignment = Alignment(horizontal="center")
        paginas.alignment = Alignment(horizontal="center")
        serie.alignment = Alignment(horizontal="center")

        # Impresora sin respuesta

        if ip.value in (None, ""):

            for celda in ws[fila]:
                celda.fill = gris

            continue

        # Color del toner

        try:

            valor = float(toner.value)

            if valor >= 60:
                toner.fill = verde

            elif valor >= 25:
                toner.fill = naranja

            else:
                toner.fill = rojo

        except Exception:
            pass

    # -----------------------------
    # Ajustar ancho
    # -----------------------------

    for columna in ws.columns:

        ancho = 0

        letra = get_column_letter(
            columna[0].column
        )

        for celda in columna:

            if celda.value is not None:

                ancho = max(
                    ancho,
                    len(str(celda.value))
                )

        ancho_columna = ancho + 3
        if str(columna[0].value).strip().upper() == "MODELO":
            ancho_columna = min(ancho_columna, 25)
        ws.column_dimensions[letra].width = ancho_columna

    # Congelar encabezado

    ws.freeze_panes = "A2"

    # Filtros

    ws.auto_filter.ref = ws.dimensions

    wb.save(ARCHIVO_SALIDA)


def actualizar():

    ultimo_json = obtener_ultimo_json()

    print(f"Usando JSON: {ultimo_json}")

    df = pd.read_excel(ARCHIVO_BASE)

    df.columns = (
        df.columns.astype(str)
        .str.strip()
        .str.upper()
    )

    df["MAC"] = df["MAC"].apply(
        normalizar_mac
    )
    ips_base = df["IP"].apply(normalizar_ip) if "IP" in df else pd.Series(index=df.index)

    datos = leer_json(ultimo_json)

    def obtener_dato(fila, campo):
        registro = buscar_impresora(
            datos,
            fila.get("MAC"),
            ips_base.loc[fila.name],
        )
        return registro.get(campo, "")

    for campo in ("MAC", "MARCA", "MODELO"):
        df[campo] = df.apply(
            lambda fila: obtener_dato(fila, campo) or fila.get(campo, ""),
            axis=1,
        )

    df["IP"] = df.apply(
        lambda fila: obtener_dato(fila, "IP") or ips_base.loc[fila.name] or "",
        axis=1,
    )

    df["NIVEL_TONER"] = df.apply(lambda fila: obtener_dato(fila, "NIVEL_TONER"), axis=1)

    df["IMPRESIONES_TOTALES"] = df.apply(
        lambda fila: obtener_dato(fila, "IMPRESIONES_TOTALES"), axis=1
    )

    df["NUMERO_SERIE"] = df.apply(lambda fila: obtener_dato(fila, "NUMERO_SERIE"), axis=1)

    df.to_excel(
        ARCHIVO_SALIDA,
        index=False,
    )

    aplicar_formato()
    try:
        exportar_pdf_libreoffice()
    except FileNotFoundError as exc:
        print(f"⚠ PDF no generado: {exc}")

    print(
        f"✔ Archivo generado: {ARCHIVO_SALIDA}"
    )




def exportar_pdf_libreoffice():

    temporal = BASE_DIR / "__tmp_impresoras.xlsx"

    shutil.copy2(ARCHIVO_SALIDA, temporal)

    wb = load_workbook(temporal)
    ws = wb.active

    # Buscar columnas por nombre
    columnas = {}

    for c in ws[1]:
        columnas[str(c.value).strip().upper()] = c.column_letter

    # Ocultar MAC e IP
    if "MAC" in columnas:
        ws.column_dimensions[columnas["MAC"]].hidden = True

    if "IP" in columnas:
        ws.column_dimensions[columnas["IP"]].hidden = True

    # Configuración de página
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    wb.save(temporal)
    wb.close()

    # Ruta de LibreOffice (soporta Linux, macOS y Windows)
    rutas = [
        shutil.which("soffice"),
        shutil.which("libreoffice"),
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    ]

    soffice = next((ruta for ruta in rutas if ruta and os.path.exists(ruta)), None)

    if soffice is None:
        raise FileNotFoundError(
            "No se encontró LibreOffice en PATH ni en la ruta por defecto.\n"
            "Instalá LibreOffice o configurá el ejecutable 'soffice'."
        )

    subprocess.run(
        [
            soffice,
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(BASE_DIR),
            str(temporal),
        ],
        check=True,
    )

    if temporal.exists():
        os.remove(temporal)

    # Renombrar el PDF sobrescribiendo cualquier archivo previo
    pdf_temporal = BASE_DIR / "__tmp_impresoras.pdf"
    pdf_salida = BASE_DIR / "impresoras_al_dia.pdf"
    if pdf_temporal.exists():
        if pdf_salida.exists():
            pdf_salida.unlink()
        pdf_temporal.rename(pdf_salida)

    print("✔ PDF generado correctamente.")

if __name__ == "__main__":
    actualizar()